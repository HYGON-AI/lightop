#!/usr/bin/env python3
"""Run and report the pytest-only cases that were blank in the old report.

Each file is run separately so the console always shows which file is being
tested and the per-file result can be saved to an Excel workbook.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime
import os
from pathlib import Path
import subprocess
import sys
import time

from category_test_runner import (
    ROOT,
    find_build_dir,
    print_case_inventory,
    run_streamed,
)


RUNNABLE_TEST_FILES = [
    "test/test_native_categorized_quant_remaining_api.py",
    "test/test_native_categorized_attention_remaining_api.py",
    "test/test_vision_existing_api.py",
    "test/test_public_envs.py",
    "test/test_native_categorized_norm_remaining_api.py",
    "test/test_native_categorized_moe_api.py",
    "test/test_kvcache_existing_api.py",
    "test/test_attention_existing_api.py",
    "test/test_native_categorized_norm_activation_api.py",
    "test/test_gemm_ops_existing_api.py",
    "test/test_quant_per_token_quant_fp8.py",
    "test/test_quant_per_token_group_quant_fp8.py",
    "test/test_moe_existing_api.py",
    "test/test_prefill_gather_and_upconvert_fp8_kv_cache.py",
    "test/test_categorized_operator_imports.py",
    "test/test_quant_blockwise_scaled_cast_transpose_api.py",
    "test/test_gemm_ops_rocblas_scaled_mm_nn.py",
    "test/test_rms_norm_dynamic_per_token_quant_sglang.py",
    "test/test_quant_destindex_copy_quantize_int8_cuda.py",
    "test/test_quant_per_token_quant_int8.py",
    "test/test_sampling_existing_api.py",
    "test/test_no_lmslim_runtime_dependency.py",
    "test/test_quant_destindex_dequantize_int8_cuda.py",
    "test/test_gemm_ops_hipblaslt_w8a8_channelwise_gemm.py",
    "test/test_quant_w4tow16.py",
    "test/test_fuse_silu_mul_clamp_quant.py",
    "test/test_activation_existing_api.py",
    "test/test_quant_fp8_to_16.py",
    "test/test_sparse_attention_transform_api.py",
    "test/test_gemm_ops_triton_scaled_mm.py",
    "test/test_per_token_group_quant_fp8_api_alignment.py",
    "test/test_loss_existing_api.py",
    "test/test_mla_qkv_a_rms_norm_dynamic_per_token_quant.py",
    "test/test_norm_existing_api.py",
    "test/test_w8a8_config_dir.py",
    "test/test_norm_lm_faster_rmsquant.py",
    "test/test_activation_lm_fuse_silu_mul_quant.py",
    "test/test_gemm_ops_triton_int8_gemm_helper.py",
    "test/test_tensor_existing_api.py",
    "test/test_native_categorized_attention_api.py",
    "test/test_gemm_ops_lazy_native_import.py",
    "test/test_native_categorized_kvcache_remaining_api.py",
    "test/test_moe_align_block_size_out.py",
    "test/test_gemm_ops_hipblaslt_w8a8_channelwise_gemm_kme.py",
    "test/test_gemm_ops_hipblaslt_w8a8_blockwise_gemm.py",
    "test/test_gemm_ops_hipblaslt_w8a8_gemm.py",
]


DEFAULT_RESULT_FILE = ROOT / "test-results" / "blank_pytest_results.xlsx"
EXCEL_CELL_LIMIT = 32_767
ILLEGAL_EXCEL_CHARACTERS = dict.fromkeys(
    (*range(0, 9), 11, 12, *range(14, 32))
)


@dataclass(frozen=True)
class TestResult:
    test_file: str
    status: str
    return_code: int
    duration_seconds: float
    started_at: datetime
    finished_at: datetime
    output: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run blank pytest cases one file at a time and save an XLSX report."
    )
    parser.add_argument(
        "--build",
        action="store_true",
        help="run `python setup.py build` before testing",
    )
    parser.add_argument(
        "--build-dir",
        type=Path,
        help="explicit build/lib.* directory to use",
    )
    parser.add_argument(
        "--max-jobs",
        type=int,
        default=16,
        help="MAX_JOBS used with --build (default: 16)",
    )
    parser.add_argument(
        "--list-cases",
        action="store_true",
        help="list runnable cases without building or testing",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_RESULT_FILE,
        help=f"XLSX result path (default: {DEFAULT_RESULT_FILE})",
    )
    return parser.parse_args()


def _excel_safe_output(output: str) -> str:
    output = output.translate(ILLEGAL_EXCEL_CHARACTERS)
    if len(output) <= EXCEL_CELL_LIMIT:
        return output

    marker = "\n\n... output truncated to fit the Excel cell limit ...\n\n"
    available = EXCEL_CELL_LIMIT - len(marker)
    head_length = available // 2
    return output[:head_length] + marker + output[-(available - head_length) :]


class XlsxResultsWriter:
    """Incrementally persist results so completed cases survive an interruption."""

    def __init__(self, output_path: Path, total_cases: int) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise SystemExit(
                "XLSX output requires openpyxl; install it with `python -m pip "
                "install openpyxl`."
            ) from exc

        self.output_path = output_path.expanduser().resolve()
        self.total_cases = total_cases
        self.results: list[TestResult] = []
        self._alignment = Alignment
        self._font = Font
        self._fill = PatternFill

        self.workbook = Workbook()
        self.sheet = self.workbook.active
        self.sheet.title = "Results"
        self._initialize_sheet()
        self._save()

    def _initialize_sheet(self) -> None:
        sheet = self.sheet
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:H1")
        sheet["A1"] = "LightOp blank pytest results"
        sheet["A1"].font = self._font(bold=True, color="FFFFFF", size=14)
        sheet["A1"].fill = self._fill("solid", fgColor="1F4E78")
        sheet["A1"].alignment = self._alignment(horizontal="center")
        sheet.row_dimensions[1].height = 24

        sheet["A2"] = "Progress"
        sheet["A2"].font = self._font(bold=True)
        sheet["B2"] = f"0/{self.total_cases} completed"

        headers = [
            "No.",
            "Test file",
            "Result",
            "Exit code",
            "Duration (s)",
            "Started at",
            "Finished at",
            "Pytest output",
        ]
        for column, value in enumerate(headers, start=1):
            cell = sheet.cell(row=4, column=column, value=value)
            cell.font = self._font(bold=True, color="FFFFFF")
            cell.fill = self._fill("solid", fgColor="4472C4")
            cell.alignment = self._alignment(horizontal="center", vertical="center")

        widths = {
            "A": 7,
            "B": 58,
            "C": 12,
            "D": 12,
            "E": 14,
            "F": 22,
            "G": 22,
            "H": 100,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A5"

    def append(self, result: TestResult) -> None:
        self.results.append(result)
        row = len(self.results) + 4
        values = [
            len(self.results),
            result.test_file,
            result.status,
            result.return_code,
            round(result.duration_seconds, 3),
            result.started_at,
            result.finished_at,
            _excel_safe_output(result.output),
        ]
        for column, value in enumerate(values, start=1):
            self.sheet.cell(row=row, column=column, value=value)

        self.sheet.cell(row=row, column=3).fill = self._fill(
            "solid", fgColor="C6EFCE" if result.status == "PASSED" else "FFC7CE"
        )
        self.sheet.cell(row=row, column=3).font = self._font(
            bold=True, color="006100" if result.status == "PASSED" else "9C0006"
        )
        self.sheet.cell(row=row, column=5).number_format = "0.000"
        for column in (6, 7):
            self.sheet.cell(row=row, column=column).number_format = "yyyy-mm-dd hh:mm:ss"
        self.sheet.cell(row=row, column=8).alignment = self._alignment(
            vertical="top", wrap_text=True
        )
        self.sheet.row_dimensions[row].height = 60
        self.sheet.auto_filter.ref = f"A4:H{row}"

        passed = sum(item.status == "PASSED" for item in self.results)
        failed = len(self.results) - passed
        self.sheet["B2"] = (
            f"{len(self.results)}/{self.total_cases} completed; "
            f"{passed} passed; {failed} failed"
        )
        self._save()

    def _save(self) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(self.output_path)


def run_test_file(test_file: str, index: int, total: int, env: dict[str, str]) -> TestResult:
    test_name = Path(test_file).name
    banner = f"testing {test_name} ({index}/{total})"
    print(f"\n{'*' * 24} {banner} {'*' * 24}\n", flush=True)

    command = [
        sys.executable,
        "-m",
        "pytest",
        "-q",
        "-rs",
        str(ROOT / test_file),
    ]
    # Excel stores local wall-clock datetimes without timezone information.
    started_at = datetime.now()
    started_clock = time.perf_counter()
    output_lines: list[str] = []
    try:
        process = subprocess.Popen(
            command,
            cwd=ROOT,
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            bufsize=1,
        )
        assert process.stdout is not None
        for line in process.stdout:
            output_lines.append(line)
            print(line, end="", flush=True)
        return_code = process.wait()
    except OSError as exc:
        return_code = -1
        message = f"Unable to start pytest: {exc}\n"
        output_lines.append(message)
        print(message, end="", flush=True)

    finished_at = datetime.now()
    duration_seconds = time.perf_counter() - started_clock
    status = "PASSED" if return_code == 0 else "FAILED"
    print(f"[{test_file}] {status} ({duration_seconds:.3f}s)", flush=True)
    return TestResult(
        test_file=test_file,
        status=status,
        return_code=return_code,
        duration_seconds=duration_seconds,
        started_at=started_at,
        finished_at=finished_at,
        output="".join(output_lines),
    )


def main() -> int:
    args = parse_args()
    if args.list_cases:
        print_case_inventory("blank-pytest", RUNNABLE_TEST_FILES, ())
        return 0

    env = os.environ.copy()
    if args.build:
        build_env = env.copy()
        build_env["MAX_JOBS"] = str(args.max_jobs)
        if not run_streamed(
            "full build", [sys.executable, "setup.py", "build"], build_env
        ):
            return 1

    build_dir = find_build_dir(args.build_dir)
    old_pythonpath = env.get("PYTHONPATH")
    env["PYTHONPATH"] = (
        str(build_dir)
        if not old_pythonpath
        else os.pathsep.join((str(build_dir), old_pythonpath))
    )
    print(f"Using LightOp build: {build_dir}")

    writer = XlsxResultsWriter(args.output, len(RUNNABLE_TEST_FILES))
    for index, test_file in enumerate(RUNNABLE_TEST_FILES, start=1):
        result = run_test_file(test_file, index, len(RUNNABLE_TEST_FILES), env)
        writer.append(result)

    print(f"\nResults saved to: {writer.output_path}")
    failed = [result for result in writer.results if result.status != "PASSED"]
    print(
        f"Summary: {len(writer.results) - len(failed)} passed, "
        f"{len(failed)} failed, {len(writer.results)} total"
    )
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
